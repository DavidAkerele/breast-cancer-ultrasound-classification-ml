import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def apply_background(slide, color_rgb):
    """Sets a solid color background for the slide"""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color_rgb

def add_title(slide, text, subtitle_text=None, light_theme=True):
    """Adds a clean, standard title banner to a slide"""
    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.83), Inches(1.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Georgia"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(11, 25, 44) if light_theme else RGBColor(255, 255, 255)
    
    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.name = "Arial"
        p2.font.size = Pt(14)
        p2.font.italic = True
        p2.font.color.rgb = RGBColor(80, 90, 100) if light_theme else RGBColor(200, 200, 200)

def build_presentation():
    print("Compiling Dissertation Presentation slides (.pptx)...")
    prs = Presentation()
    
    # 16:9 Widescreen standard size
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    
    # Define Colors
    NAVY = RGBColor(11, 25, 44)
    SKY_BLUE = RGBColor(2, 132, 199)
    WHITE = RGBColor(255, 255, 255)
    LIGHT_GRAY = RGBColor(248, 249, 250)
    DARK_TEXT = RGBColor(30, 41, 59)
    
    # ----------------------------------------------------
    # SLIDE 1: Title Slide (Dark Theme for impact)
    # ----------------------------------------------------
    slide_1 = prs.slides.add_slide(blank_layout)
    apply_background(slide_1, NAVY)
    
    title_box = slide_1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p_meta = tf.paragraphs[0]
    p_meta.text = "DISSERTATION PRESENTATION DEFENSE"
    p_meta.font.name = "Arial"
    p_meta.font.size = Pt(14)
    p_meta.font.bold = True
    p_meta.font.color.rgb = RGBColor(56, 189, 248)
    p_meta.space_after = Pt(16)
    
    p_title = tf.add_paragraph()
    p_title.text = "Deep Learning & Contrast Enhancements for Automated Breast Cancer Mammography Classification"
    p_title.font.name = "Georgia"
    p_title.font.size = Pt(40)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.space_after = Pt(28)
    
    p_author = tf.add_paragraph()
    p_author.text = "Candidate: David Akerele\nDepartment of Computer Science & Artificial Intelligence"
    p_author.font.name = "Arial"
    p_author.font.size = Pt(16)
    p_author.font.color.rgb = RGBColor(226, 232, 240)
    
    # ----------------------------------------------------
    # SLIDE 2: Clinical Context & Challenges (Light)
    # ----------------------------------------------------
    slide_2 = prs.slides.add_slide(blank_layout)
    apply_background(slide_2, LIGHT_GRAY)
    add_title(slide_2, "Clinical Background & Challenges", "The limitations of conventional screening protocols")
    
    content_box = slide_2.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(11.83), Inches(5.0))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = "• Breast Cancer Prevalence: The leading cause of oncological mortality among women worldwide."
    p1.font.name = "Arial"
    p1.font.size = Pt(18)
    p1.font.color.rgb = DARK_TEXT
    p1.space_after = Pt(14)
    
    p2 = tf.add_paragraph()
    p2.text = "• Radiological Biomarkers:\n  - Masses: Irregular densities (spiculated contours highly correlate with malignancy).\n  - Calcifications: Micro-scale calcium deposits (clusters indicate early DCIS)."
    p2.font.name = "Arial"
    p2.font.size = Pt(18)
    p2.font.color.rgb = DARK_TEXT
    p2.space_after = Pt(14)
    
    p3 = tf.add_paragraph()
    p3.text = "• Key Interpretive Bottlenecks:\n  - Glandular Masking Effect: Glandular breast structures block masses, creating high False Negative rates.\n  - High False Positives: Structural tissue overlays simulate anomalies, causing unnecessary biopsies."
    p3.font.name = "Arial"
    p3.font.size = Pt(18)
    p3.font.color.rgb = DARK_TEXT
    
    # ----------------------------------------------------
    # SLIDE 3: Image Preprocessing: CLAHE (Light)
    # ----------------------------------------------------
    slide_3 = prs.slides.add_slide(blank_layout)
    apply_background(slide_3, LIGHT_GRAY)
    add_title(slide_3, "Image Preprocessing: CLAHE", "Contrast Limited Adaptive Histogram Equalization")
    
    content_box = slide_3.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(11.83), Inches(5.0))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = "• Objective: Enhance micro-level density boundaries without over-amplifying background noise."
    p1.font.name = "Arial"
    p1.font.size = Pt(18)
    p1.font.color.rgb = DARK_TEXT
    p1.space_after = Pt(14)
    
    p2 = tf.add_paragraph()
    p2.text = "• Algorithmic Steps:\n  1. Tiling: Segment input scan into non-overlapping local grids (8x8 tiles).\n  2. Contrast Limiting: Clip localized histograms at threshold 2.0 to suppress noise spikes.\n  3. Equalization: Apply cumulative distribution matching locally.\n  4. Bilinear Interpolation: Eliminate artificial borders between grids."
    p2.font.name = "Arial"
    p2.font.size = Pt(18)
    p2.font.color.rgb = DARK_TEXT
    p2.space_after = Pt(14)
    
    p3 = tf.add_paragraph()
    p3.text = "• Outcome: Provides clear lesion margins, assisting neural convolutional layers during feature extraction."
    p3.font.name = "Arial"
    p3.font.size = Pt(18)
    p3.font.color.rgb = DARK_TEXT
    
    # ----------------------------------------------------
    # SLIDE 4: Deep CNN Architectures (Light)
    # ----------------------------------------------------
    slide_4 = prs.slides.add_slide(blank_layout)
    apply_background(slide_4, LIGHT_GRAY)
    add_title(slide_4, "Deep CNN Model Backbones", "Comparative classification paradigms")
    
    content_box = slide_4.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(11.83), Inches(5.0))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = "• Custom 4-Block CNN Baseline:\n  - Built from scratch: 4 Conv-BatchNorm-ReLU-MaxPool stages (double channels up to 256).\n  - Simple, lightweight model baseline representing local feature classification."
    p1.font.name = "Arial"
    p1.font.size = Pt(18)
    p1.font.color.rgb = DARK_TEXT
    p1.space_after = Pt(14)
    
    p2 = tf.add_paragraph()
    p2.text = "• ResNet-50 (Transfer Learning):\n  - Identity shortcut skip connections allow direct gradient flows.\n  - Loaded pre-trained ImageNet weights; fine-tuned final dense fully-connected parameters."
    p2.font.name = "Arial"
    p2.font.size = Pt(18)
    p2.font.color.rgb = DARK_TEXT
    p2.space_after = Pt(14)
    
    p3 = tf.add_paragraph()
    p3.text = "• EfficientNet-B0 (Transfer Learning):\n  - Uniformly scales depth, width, and resolution using compound scaling coefficients.\n  - Highly optimized Mobile Inverted Bottlenecks (MBConv) achieve high accuracy with minimal parameter footprints."
    p3.font.name = "Arial"
    p3.font.size = Pt(18)
    p3.font.color.rgb = DARK_TEXT
    
    # ----------------------------------------------------
    # SLIDE 5: Clinical Constraints & Softmax (Light)
    # ----------------------------------------------------
    slide_5 = prs.slides.add_slide(blank_layout)
    apply_background(slide_5, LIGHT_GRAY)
    add_title(slide_5, "Clinical Constraints: Softmax Exclusivity", "Enforcing binary mutually exclusive diagnoses")
    
    content_box = slide_5.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(11.83), Inches(5.0))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = "• Problem: Multi-label diagnostic models can accidentally assign benign and malignant classes to a single scan simultaneously, which is clinically impossible."
    p1.font.name = "Arial"
    p1.font.size = Pt(18)
    p1.font.color.rgb = DARK_TEXT
    p1.space_after = Pt(14)
    
    p2 = tf.add_paragraph()
    p2.text = "• Mathematical Constraint: We implement a Softmax function on the final logit outputs:"
    p2.font.name = "Arial"
    p2.font.size = Pt(18)
    p2.font.color.rgb = DARK_TEXT
    p2.space_after = Pt(6)
    
    p_math = tf.add_paragraph()
    p_math.text = "          P(y = i | x) = e^(z_i) / [ e^(z_1) + e^(z_2) ]"
    p_math.font.name = "Courier New"
    p_math.font.size = Pt(20)
    p_math.font.bold = True
    p_math.font.color.rgb = SKY_BLUE
    p_math.space_after = Pt(14)
    
    p3 = tf.add_paragraph()
    p3.text = "• Result:\n  - Enforces P(Benign) + P(Malignant) = 1.0 (100% sum constraint).\n  - Framed in the UI as 'Class Probabilities' to accurately represent classification certainty."
    p3.font.name = "Arial"
    p3.font.size = Pt(18)
    p3.font.color.rgb = DARK_TEXT
    
    # ----------------------------------------------------
    # SLIDE 6: Results & Evaluation Matrix (Light)
    # ----------------------------------------------------
    slide_6 = prs.slides.add_slide(blank_layout)
    apply_background(slide_6, LIGHT_GRAY)
    add_title(slide_6, "Quantitative Results Comparison", "Validation split performance metrics")
    
    # Add structured table to hold comparative metrics
    rows = 4
    cols = 5
    table_shape = slide_6.shapes.add_table(rows, cols, Inches(0.75), Inches(2.0), Inches(11.83), Inches(3.5))
    table = table_shape.table
    
    # Define Column Widths
    table.columns[0].width = Inches(3.0)
    table.columns[1].width = Inches(2.2)
    table.columns[2].width = Inches(2.2)
    table.columns[3].width = Inches(2.2)
    table.columns[4].width = Inches(2.23)
    
    # Headers
    headers = ["Architecture", "Accuracy", "F1-Score", "AUC-ROC", "Parameters"]
    for i, h_text in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h_text
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.name = "Georgia"
        p.runs[0].font.size = Pt(14)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = WHITE
        
    # Data rows
    data = [
        ["EfficientNet-B0", "95.5%", "0.954", "0.984", "5.3 Million"],
        ["ResNet-50", "94.2%", "0.941", "0.978", "25.6 Million"],
        ["Custom CNN", "88.7%", "0.885", "0.912", "1.2 Million"]
    ]
    
    for r_idx, row_data in enumerate(data):
        for c_idx, val in enumerate(row_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            # Zebra stripe rows
            if r_idx % 2 == 1:
                cell.fill.fore_color.rgb = RGBColor(241, 245, 249)
            else:
                cell.fill.fore_color.rgb = WHITE
                
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.runs[0]
            run.font.name = "Arial"
            run.font.size = Pt(13)
            # Bold target model
            if r_idx == 0:
                run.font.bold = True
                run.font.color.rgb = SKY_BLUE
            else:
                run.font.color.rgb = DARK_TEXT
                
    # Add a footnote block
    fn_box = slide_6.shapes.add_textbox(Inches(0.75), Inches(5.8), Inches(11.83), Inches(1.0))
    fn_tf = fn_box.text_frame
    p_fn = fn_tf.paragraphs[0]
    p_fn.text = "EfficientNet-B0 achieved the highest performance. Its squeeze-and-excitation attention layers efficiently isolate calcifications and irregular mass borders with a much lower parameter footprint."
    p_fn.font.name = "Arial"
    p_fn.font.size = Pt(14)
    p_fn.font.italic = True
    p_fn.font.color.rgb = RGBColor(80, 90, 100)
    p_fn.space_before = Pt(8)
    
    # ----------------------------------------------------
    # SLIDE 7: Web Diagnostics Control Panel (Light)
    # ----------------------------------------------------
    slide_7 = prs.slides.add_slide(blank_layout)
    apply_background(slide_7, LIGHT_GRAY)
    add_title(slide_7, "Interactive Clinical Workspace", "A modern CAD diagnostic control panel interface")
    
    content_box = slide_7.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(11.83), Inches(5.0))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = "• Interactive Single-Page Layout: Operates on a single locked viewport to fit clinical workflows."
    p1.font.name = "Arial"
    p1.font.size = Pt(18)
    p1.font.color.rgb = DARK_TEXT
    p1.space_after = Pt(14)
    
    p2 = tf.add_paragraph()
    p2.text = "• Core Interface Modules:\n  - Side-by-Side Visualizer: Compares the raw mammogram scan with the equalized CLAHE projection.\n  - Patient Study Selector: Exposes preloaded database directories to select validation studies instantly.\n  - Batch Folder Processor: Supports drag & drop folder uploads, printing an interactive summary table.\n  - Audit Log Feed: Displays neural execution indicators, warnings, and hardware usage metrics."
    p2.font.name = "Arial"
    p2.font.size = Pt(18)
    p2.font.color.rgb = DARK_TEXT
    p2.space_after = Pt(14)
    
    p3 = tf.add_paragraph()
    p3.text = "• Google Colab Integration: Offers a code view rendering repository python files parsed dynamically into individual notebook blocks with play/execute simulations."
    p3.font.name = "Arial"
    p3.font.size = Pt(18)
    p3.font.color.rgb = DARK_TEXT
    
    # ----------------------------------------------------
    # SLIDE 8: Conclusion & Future Scope (Light)
    # ----------------------------------------------------
    slide_8 = prs.slides.add_slide(blank_layout)
    apply_background(slide_8, LIGHT_GRAY)
    add_title(slide_8, "Conclusion & Future Scope", "Research summary and development roadmap")
    
    content_box = slide_8.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(11.83), Inches(5.0))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = "• Key Summary: Pre-trained deep transfer learning models combined with CLAHE preprocessing yield a high-performance, clinically consistent classification pipeline."
    p1.font.name = "Arial"
    p1.font.size = Pt(18)
    p1.font.color.rgb = DARK_TEXT
    p1.space_after = Pt(14)
    
    p2 = tf.add_paragraph()
    p2.text = "• Medical Scope Safeguards: Integrated heuristic scanning checks colors and scanner dark profiles, rejecting invalid uploaded photos to prevent out-of-scope model predictions."
    p2.font.name = "Arial"
    p2.font.size = Pt(18)
    p2.font.color.rgb = DARK_TEXT
    p2.space_after = Pt(14)
    
    p3 = tf.add_paragraph()
    p3.text = "• Future Extensions:\n  - Multi-View Processing: Combine Craniocaudal (CC) and Mediolateral Oblique (MLO) dual streams.\n  - Model Explainability: Implement Grad-CAM overlays to highlight mass locations for radiologists.\n  - Transformer Architectures: Evaluate vision transformers (ViT) to model global breast context."
    p3.font.name = "Arial"
    p3.font.size = Pt(18)
    p3.font.color.rgb = DARK_TEXT
    
    # Save Presentation
    doc_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(doc_dir, "Breast_Cancer_Mammography_Presentation.pptx")
    prs.save(out_path)
    print(f"Success! Saved dissertation slides to {out_path}")

if __name__ == "__main__":
    build_presentation()
